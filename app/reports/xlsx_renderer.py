import io
from datetime import datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
WHITE_FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)
REGULAR_FONT = Font(name="Calibri", size=11)
MUTED_FONT = Font(name="Calibri", size=9, italic=True, color="6B7280")

THIN_SIDE = Side(border_style="thin", color="E5E7EB")
BORDER_BOX = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def _format_cell_value(val: Any) -> Any:
    if isinstance(val, Decimal):
        return float(val)
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(val, list):
        return ", ".join(str(x) for x in val)
    return val


def _auto_fit_columns(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def _add_sheet_header(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    title: str,
    subtitle: Optional[str] = None,
    filters_desc: Optional[str] = None,
) -> int:
    ws.append([title])
    ws.cell(row=1, column=1).font = TITLE_FONT
    current_row = 2

    if subtitle:
        ws.append([subtitle])
        ws.cell(row=current_row, column=1).font = MUTED_FONT
        current_row += 1

    if filters_desc:
        ws.append([f"Filters: {filters_desc}"])
        ws.cell(row=current_row, column=1).font = MUTED_FONT
        current_row += 1

    ws.append([])  # blank row
    return current_row + 1


def _write_table(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    start_row: int,
    headers: List[str],
    rows: List[List[Any]],
) -> None:
    # Write header
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=c_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT_BOLD
        cell.alignment = Alignment(horizontal="center" if "id" in h.lower() or "%" in h or "date" in h.lower() else "left")
        cell.border = BORDER_BOX

    # Write data rows
    for r_offset, row_data in enumerate(rows):
        r_idx = start_row + 1 + r_offset
        is_zebra = r_offset % 2 == 1
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=_format_cell_value(val))
            cell.font = REGULAR_FONT
            cell.border = BORDER_BOX
            if is_zebra:
                cell.fill = ZEBRA_FILL
            if isinstance(val, (int, float, Decimal)):
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, (datetime,)):
                cell.alignment = Alignment(horizontal="center")


def render_report_to_xlsx(report_type: str, data: Dict[str, Any], filters: Dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    filter_summary = f"Period: {filters.get('period', 'N/A')}"
    if filters.get("team_id"):
        filter_summary += f" | Team: {filters['team_id']}"
    if filters.get("rep_id"):
        filter_summary += f" | Rep: {filters['rep_id']}"
    if filters.get("approval_status"):
        filter_summary += f" | Status: {filters['approval_status']}"

    if report_type == "summary":
        ws = wb.create_sheet(title="Executive Summary")
        start = _add_sheet_header(ws, "DealFlow360 - Executive Summary Report", "Governance & Sales Operations Platform", filter_summary)
        headers = ["Metric", "Value"]
        rows = [
            ["Quotations Created", data.get("quotations_created", 0)],
            ["Quotations Sent", data.get("quotations_sent", 0)],
            ["Quotations Confirmed", data.get("quotations_confirmed", 0)],
            ["Win Rate (%)", f"{float(data.get('win_rate', 0.0)):.2f}%"],
            ["Confirmed Revenue", f"{float(data.get('revenue_confirmed', 0.0)):,.2f}"],
            ["Average Discount (%)", f"{float(data.get('avg_discount', 0.0)):.2f}%"],
            ["Average Margin (%)", f"{float(data.get('avg_margin', 0.0)):.2f}%"],
            ["Avg Approval Turnaround (Hours)", f"{float(data.get('avg_approval_turnaround_hours', 0.0)):.2f}"],
            ["Pending Approvals Count", data.get("pending_approvals_count", 0)],
            ["Generated At", data.get("generated_at", "")],
        ]
        _write_table(ws, start, headers, rows)
        _auto_fit_columns(ws)

    elif report_type in ("deals", "quotations"):
        ws = wb.create_sheet(title="Deals")
        start = _add_sheet_header(ws, "DealFlow360 - Deals & Quotations Report", "Governance & Pipeline Visibility", filter_summary)
        headers = [
            "Reference", "Odoo Order", "Partner Name", "Owner UID", "Status",
            "Approval State", "Health", "Risk Score", "Amount Total", "Discount %",
            "Margin %", "Created At", "Confirmed At"
        ]
        items = data.get("items", [])
        rows = []
        for it in items:
            rows.append([
                it.get("reference", ""),
                it.get("odoo_order_name", ""),
                it.get("partner_name", ""),
                it.get("owner_odoo_user_id", ""),
                it.get("status", ""),
                it.get("approval_state", ""),
                it.get("health_status", ""),
                float(it.get("current_risk_score", 0.0)) if it.get("current_risk_score") is not None else "N/A",
                float(it.get("amount_total", 0.0)),
                float(it.get("order_discount_pct", 0.0)),
                float(it.get("margin_pct", 0.0)),
                it.get("created_at", ""),
                it.get("confirmed_at", ""),
            ])
        _write_table(ws, start, headers, rows)
        _auto_fit_columns(ws)

    elif report_type == "approvals":
        # Sheet 1: Turnaround by Stage
        ws1 = wb.create_sheet(title="Turnaround By Stage")
        start1 = _add_sheet_header(ws1, "Approvals Turnaround by Stage", filters_desc=filter_summary)
        h1 = ["Stage / Level", "Total Requests", "Avg Turnaround (Hours)", "Approved", "Rejected", "Returned"]
        r1 = []
        for st in data.get("turnaround_by_stage", []):
            r1.append([
                st.get("required_level", ""),
                st.get("total_requests", 0),
                float(st.get("avg_turnaround_hours", 0.0)),
                st.get("approved_count", 0),
                st.get("rejected_count", 0),
                st.get("returned_count", 0),
            ])
        _write_table(ws1, start1, h1, r1)
        _auto_fit_columns(ws1)

        # Sheet 2: Status Breakdown
        ws2 = wb.create_sheet(title="Status Breakdown")
        start2 = _add_sheet_header(ws2, "Approval Requests by Status", filters_desc=filter_summary)
        h2 = ["Status", "Count"]
        r2 = [[x.get("status", ""), x.get("count", 0)] for x in data.get("by_status", [])]
        _write_table(ws2, start2, h2, r2)
        _auto_fit_columns(ws2)

        # Sheet 3: Top Rejection Reasons
        ws3 = wb.create_sheet(title="Rejection Reasons")
        start3 = _add_sheet_header(ws3, "Top Rejection & Return Reasons", filters_desc=filter_summary)
        h3 = ["Reason", "Occurrences"]
        r3 = [[x.get("reason", ""), x.get("count", 0)] for x in data.get("top_rejection_reasons", [])]
        _write_table(ws3, start3, h3, r3)
        _auto_fit_columns(ws3)

    elif report_type == "products":
        # Sheet 1: Best Selling by Qty
        ws1 = wb.create_sheet(title="Best Selling (Qty)")
        start1 = _add_sheet_header(ws1, "Best Selling Products by Quantity", filters_desc=filter_summary)
        h = ["Product ID", "Product Name", "Category ID", "Qty Sold", "Revenue", "Avg Discount %"]
        r1 = []
        for p in data.get("best_selling_by_qty", []):
            r1.append([
                p.get("product_id"), p.get("product_name"), p.get("category_id") or "N/A",
                p.get("qty_sold"), float(p.get("revenue", 0.0)), float(p.get("avg_discount_pct", 0.0))
            ])
        _write_table(ws1, start1, h, r1)
        _auto_fit_columns(ws1)

        # Sheet 2: Best Selling by Revenue
        ws2 = wb.create_sheet(title="Best Selling (Revenue)")
        start2 = _add_sheet_header(ws2, "Best Selling Products by Revenue", filters_desc=filter_summary)
        r2 = []
        for p in data.get("best_selling_by_revenue", []):
            r2.append([
                p.get("product_id"), p.get("product_name"), p.get("category_id") or "N/A",
                p.get("qty_sold"), float(p.get("revenue", 0.0)), float(p.get("avg_discount_pct", 0.0))
            ])
        _write_table(ws2, start2, h, r2)
        _auto_fit_columns(ws2)

        # Sheet 3: Most Discounted
        ws3 = wb.create_sheet(title="Most Discounted")
        start3 = _add_sheet_header(ws3, "Most Discounted Products", filters_desc=filter_summary)
        r3 = []
        for p in data.get("most_discounted", []):
            r3.append([
                p.get("product_id"), p.get("product_name"), p.get("category_id") or "N/A",
                p.get("qty_sold"), float(p.get("revenue", 0.0)), float(p.get("avg_discount_pct", 0.0))
            ])
        _write_table(ws3, start3, h, r3)
        _auto_fit_columns(ws3)

    elif report_type == "discounts":
        ws1 = wb.create_sheet(title="Rep Discount Exposure")
        start1 = _add_sheet_header(ws1, "Rep Discount Exposure & Anomalies", filters_desc=filter_summary)
        h1 = [
            "Rep ID", "Rep Name", "Team ID", "Deals Count", "Total Revenue",
            "Avg Weighted Discount %", "Overage Count", "Overage Freq %", "Anomalies", "Max Discount %"
        ]
        r1 = []
        for s in data.get("rep_stats", []):
            r1.append([
                s.get("rep_id"), s.get("rep_name"), s.get("team_id") or "N/A",
                s.get("deals_count"), float(s.get("total_revenue", 0.0)),
                float(s.get("avg_weighted_discount_pct", 0.0)),
                s.get("overage_count"), float(s.get("overage_frequency_pct", 0.0)),
                s.get("anomaly_count"), float(s.get("max_discount_pct", 0.0))
            ])
        _write_table(ws1, start1, h1, r1)
        _auto_fit_columns(ws1)

        ws2 = wb.create_sheet(title="Company Summary")
        start2 = _add_sheet_header(ws2, "Company Discount Metrics", filters_desc=filter_summary)
        h2 = ["Metric", "Value"]
        r2 = [
            ["Company Avg Discount (%)", f"{float(data.get('company_avg_discount_pct', 0.0)):.2f}%"],
            ["Total Discount Anomalies", data.get("total_anomalies", 0)],
            ["Total Policy Overages", data.get("total_overages", 0)],
        ]
        _write_table(ws2, start2, h2, r2)
        _auto_fit_columns(ws2)

    elif report_type in ("pipeline", "risk"):
        ws1 = wb.create_sheet(title="Pipeline Funnel")
        start1 = _add_sheet_header(ws1, "Pipeline Funnel per Stage", filters_desc=filter_summary)
        h1 = ["Stage / Status", "Deals Count", "Total Value", "Avg Discount %", "Avg Margin %"]
        r1 = []
        for p in data.get("pipeline_funnel", []):
            r1.append([
                p.get("status"), p.get("count"), float(p.get("total_value", 0.0)),
                float(p.get("avg_discount_pct", 0.0)), float(p.get("avg_margin_pct", 0.0))
            ])
        _write_table(ws1, start1, h1, r1)
        _auto_fit_columns(ws1)

        ws2 = wb.create_sheet(title="Risk Distribution")
        start2 = _add_sheet_header(ws2, "Risk Score Distribution", filters_desc=filter_summary)
        h2 = ["Severity", "Deals Count", "Total Value", "Avg Risk Score"]
        r2 = []
        for r in data.get("risk_distribution", []):
            r2.append([
                r.get("severity"), r.get("count"), float(r.get("total_value", 0.0)),
                float(r.get("avg_risk_score", 0.0))
            ])
        _write_table(ws2, start2, h2, r2)
        _auto_fit_columns(ws2)

        ws3 = wb.create_sheet(title="Top Risk Factors")
        start3 = _add_sheet_header(ws3, "Top Risk Factors Triggered", filters_desc=filter_summary)
        h3 = ["Risk Factor Type", "Count", "Avg Contribution"]
        r3 = []
        for f in data.get("top_risk_factors", []):
            r3.append([f.get("factor_type"), f.get("count"), float(f.get("avg_contribution", 0.0))])
        _write_table(ws3, start3, h3, r3)
        _auto_fit_columns(ws3)

    elif report_type == "fulfillment":
        ws1 = wb.create_sheet(title="Warehouse Breakdown")
        start1 = _add_sheet_header(ws1, "Fulfillment per Warehouse", filters_desc=filter_summary)
        h1 = ["Warehouse ID", "Warehouse Name", "Shipments", "Allocated Qty", "Shipping Cost"]
        r1 = []
        for w in data.get("warehouse_breakdown", []):
            r1.append([
                w.get("warehouse_id"), w.get("warehouse_name"), w.get("shipment_count"),
                w.get("total_allocated_qty"), float(w.get("total_shipping_cost", 0.0))
            ])
        _write_table(ws1, start1, h1, r1)
        _auto_fit_columns(ws1)

        ws2 = wb.create_sheet(title="Fulfillment KPIs")
        start2 = _add_sheet_header(ws2, "Fulfillment Efficiency KPIs", filters_desc=filter_summary)
        h2 = ["Metric", "Value"]
        r2 = [
            ["Multi-Warehouse Split Rate (%)", f"{float(data.get('split_rate', 0.0)):.2f}%"],
            ["Backorder Rate (%)", f"{float(data.get('backorder_rate', 0.0)):.2f}%"],
            ["On-Time Fulfillment (%)", f"{float(data.get('on_time_pct', 0.0)):.2f}%"],
            ["Total Fulfillment Plans", data.get("total_plans", 0)],
            ["Total Shipments", data.get("total_shipments", 0)],
        ]
        _write_table(ws2, start2, h2, r2)
        _auto_fit_columns(ws2)

    elif report_type == "billing":
        ws1 = wb.create_sheet(title="Invoiced vs Paid")
        start1 = _add_sheet_header(ws1, "Billing: Invoiced vs Paid & Overdue", filters_desc=filter_summary)
        h1 = ["Metric", "Value"]
        r1 = [
            ["Total Invoiced Amount", f"{float(data.get('total_invoiced', 0.0)):,.2f}"],
            ["Total Amount Paid", f"{float(data.get('total_paid', 0.0)):,.2f}"],
            ["Outstanding Balance", f"{float(data.get('total_outstanding', 0.0)):,.2f}"],
            ["Overdue Invoices Amount", f"{float(data.get('overdue_amount', 0.0)):,.2f}"],
            ["Overdue Invoices Count", data.get("overdue_count", 0)],
        ]
        _write_table(ws1, start1, h1, r1)
        _auto_fit_columns(ws1)

        ws2 = wb.create_sheet(title="Subscriptions & Credits")
        start2 = _add_sheet_header(ws2, "Subscriptions & Credit Notes", filters_desc=filter_summary)
        h2 = ["Metric", "Value"]
        r2 = [
            ["Monthly Recurring Revenue (MRR)", f"{float(data.get('mrr', 0.0)):,.2f}"],
            ["Active Subscriptions Count", data.get("active_subscriptions_count", 0)],
            ["Total Credit Notes Amount", f"{float(data.get('total_credit_notes_amount', 0.0)):,.2f}"],
            ["Credit Notes Count", data.get("credit_notes_count", 0)],
        ]
        _write_table(ws2, start2, h2, r2)
        _auto_fit_columns(ws2)

    else:
        # Fallback generic sheet
        ws = wb.create_sheet(title="Report")
        start = _add_sheet_header(ws, f"DealFlow360 - {report_type.title()} Report", filters_desc=filter_summary)
        h = ["Key", "Value"]
        r = [[k, str(v)] for k, v in data.items() if not isinstance(v, (list, dict))]
        _write_table(ws, start, h, r)
        _auto_fit_columns(ws)

    # Remove default initial sheet
    if default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
